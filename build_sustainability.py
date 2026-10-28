"""Build the separate, offline S1/S2/S3 dashboard without touching the SRS dashboard."""
import argparse,datetime,json,pathlib,hashlib
import openpyxl
ROOT=pathlib.Path(__file__).parent

def extract_sustainability(path):
    cached=openpyxl.load_workbook(path,data_only=True)
    formulas=openpyxl.load_workbook(path,data_only=False)
    names={}
    for name in cached.sheetnames:
        key=name.strip().casefold()
        if key in names:raise ValueError('Dubbele tabbladnamen na normalisatie: '+name)
        names[key]=name
    result={}
    for name,start,width in [('S1',7,7),('S2',8,6),('S3',12,9)]:
        actual=names.get(name.lower())
        if actual is None:raise ValueError('Tabblad ontbreekt: '+name)
        rows=[]
        for cells in cached[actual].iter_rows(min_row=start,max_col=width):
            values=[]
            for cell in cells:
                value=cell.value
                if isinstance(value,(datetime.datetime,datetime.date,datetime.time)):value=value.isoformat()
                if isinstance(value,datetime.timedelta):value=value.total_seconds()/86400
                if cell.data_type=='e':value={'error':value}
                if value is None and formulas[actual][cell.coordinate].data_type=='f':value={'error':'Formule-uitkomst ontbreekt'}
                values.append(value)
            if not any(x is not None for x in values):continue
            if name=='S1' and str(values[0]).strip().casefold() in ('naam','naam leverancier'):continue
            rows.append({'row':cells[0].row,'v':values})
        result[name]=rows
    return result

if __name__=='__main__':
    parser=argparse.ArgumentParser();parser.add_argument('excel');args=parser.parse_args()
    data=extract_sustainability(args.excel)
    html=(ROOT/'sustainability.template.html').read_text().replace('/* SUSTAINABILITY_APP */',(ROOT/'sustainability.js').read_text()).replace('/* EMBEDDED_SUSTAINABILITY */',json.dumps(data,ensure_ascii=False).replace('</','<\\/')).replace('/* EXCEL_PARSER */',(ROOT/'vendor/xlsx.full.min.js').read_text())
    output=ROOT/'Alliander_SRS_Duurzaamheid.html';output.write_text(html)
    print(json.dumps({'output':output.name,'records':{k:len(v) for k,v in data.items()},'source_sha256':hashlib.sha256(pathlib.Path(args.excel).read_bytes()).hexdigest()},ensure_ascii=False))
